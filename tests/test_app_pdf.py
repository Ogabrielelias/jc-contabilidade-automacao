import io
import tempfile
import unittest
from pathlib import Path

import openpyxl
from streamlit.testing.v1 import AppTest

from folha_config import CHAVE_CONFIG_PDF, COLUNAS_PDF, config_pdf_inicial

ROOT = Path(__file__).resolve().parents[1]


@unittest.skipUnless((ROOT / "Buffon Folha de salário 09 2026.pdf").exists(), "PDF de exemplo ausente")
class AppPdfTest(unittest.TestCase):
    def test_formulas_recalculam_salarios_e_totais(self):
        from tests.support.calcular_excel import calcular
        app = self.iniciar()
        app.session_state['test_provisao'] = True
        app.session_state['test_inss'] = True
        app.session_state['test_adiantamento'] = True
        app.run()
        app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        arquivo = app.session_state['test_export']
        workbook = openpyxl.load_workbook(io.BytesIO(arquivo))
        cache = openpyxl.load_workbook(io.BytesIO(arquivo), data_only=True)
        ler = calcular(workbook)
        quantidade = 0
        for aba in workbook:
            for row in aba:
                for celula in row:
                    if celula.data_type != 'f':
                        continue
                    self.assertNotIn('Base Folha', celula.value)
                    self.assertNotIn('Base Provisão', celula.value)
                    self.assertNotIn('SUMIFS(', celula.value)
                    self.assertLessEqual(len(celula.value), 8192)
                    esperado = cache[aba.title][celula.coordinate].value
                    atual = ler(aba.title, celula.coordinate)
                    if esperado is None:
                        self.assertEqual(atual, '', (aba.title, celula.coordinate))
                    else:
                        self.assertAlmostEqual(atual, esperado, places=2, msg=f'{aba.title}!{celula.coordinate}')
                    quantidade += 1
        self.assertGreater(quantidade, 1000)
        self.assertNotIn('Base Folha', workbook.sheetnames)
        self.assertNotIn('Base Provisão', workbook.sheetnames)
        # Os valores de origem estão embutidos nas fórmulas.
        ws = workbook['Resumo Salários']
        coluna_resultado = next(c.column_letter for c in ws[5] if c.value == 'Resultado')
        antes = ler(ws.title, coluna_resultado+'6')
        coluna_salario = next(c.column_letter for c in ws[5] if c.value == 'T. salário')
        ws[coluna_salario+'6'] = ws[coluna_salario+'6'].value + '+100'
        self.assertAlmostEqual(calcular(workbook)(ws.title, coluna_resultado+'6'), antes+100, places=2)
        antes = ler('Provisão 13º', 'I11')
        workbook['Provisão 13º']['E11'] = workbook['Provisão 13º']['E11'].value + '+50'
        self.assertAlmostEqual(calcular(workbook)('Provisão 13º', 'I11'), antes+50, places=2)

    def test_provisao_sozinha_e_no_excel_geral(self):
        app = self.iniciar()
        app.session_state['test_provisao'] = True
        app.session_state['test_somente_provisao'] = True
        app.run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertTrue(any('Filial 119' in w.value for w in app.warning))
        self.assertEqual(app.tabs[0].label, 'Provisão de 13º salário')
        self.assertEqual(len(app.tabs[0].dataframe), 1)
        workbook = openpyxl.load_workbook(io.BytesIO(app.session_state['test_export']), data_only=True)
        self.assertEqual(workbook.sheetnames, ['Provisão 13º'])
        self.assertEqual(workbook.active['I11'].value, 130083.78)
        app.session_state['test_somente_provisao'] = False
        app.run()
        app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        workbook = openpyxl.load_workbook(io.BytesIO(app.session_state['test_export']))
        self.assertIn('Resumo Salários', workbook.sheetnames)
        self.assertIn('Provisão 13º', workbook.sheetnames)
        self.assertEqual(app.tabs[6].label, 'Provisão de 13º salário')
        self.assertEqual(len(app.tabs[6].dataframe), 1)
        app.session_state['test_provisao_incompleta'] = True
        app.run()
        self.assertFalse(app.get('download_button'))

    def test_inss_sozinho_e_no_excel_geral(self):
        app = self.iniciar()
        app.session_state['test_inss'] = True
        app.session_state['test_somente_inss'] = True
        app.run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        workbook = openpyxl.load_workbook(io.BytesIO(app.session_state['test_export']))
        self.assertEqual(workbook.sheetnames, ['Resumo INSS'])
        self.assertEqual(workbook.active['D5'].value, '=B5+C5')
        app.session_state['test_somente_inss'] = False
        app.run()
        app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        workbook = openpyxl.load_workbook(io.BytesIO(app.session_state['test_export']))
        self.assertIn('Resumo Salários', workbook.sheetnames)
        self.assertIn('Resumo INSS', workbook.sheetnames)
        self.assertEqual(workbook['Resumo INSS']['D121'].value, '=SUM(D5:D120)')
        app.session_state['test_inss_incompleto'] = True
        app.run()
        self.assertFalse(app.exception)
        self.assertTrue(any('dois arquivos' in aviso.value for aviso in app.warning))
        self.assertFalse(app.get('download_button'))

    def test_entrada_13_rejeita_modelo_antigo(self):
        app = self.iniciar()
        app.session_state["test_pdf_13_invalido"] = True
        app.run()
        self.assertFalse(app.exception)
        self.assertTrue(app.error)
        self.assertIn("O modelo antigo de provisão não é aceito", app.error[0].value)
        self.assertNotIn("test_export", app.session_state)

    def test_adiantamento_sozinho_e_com_folha(self):
        app = self.iniciar()
        app.session_state["test_adiantamento"] = True
        app.session_state["test_somente_adiantamento"] = True
        app.run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        workbook = openpyxl.load_workbook(io.BytesIO(app.session_state["test_export"]))
        self.assertEqual(workbook.sheetnames, ["Ad 13 Sal e FGTS"])
        self.assertEqual(
            app.session_state["test_export_name"],
            "resumo_adiantamento_13_buffon_09-2026.xlsx",
        )
        app.session_state["test_somente_adiantamento"] = False
        app.run()
        app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        workbook = openpyxl.load_workbook(io.BytesIO(app.session_state["test_export"]))
        self.assertIn("Resumo Salários", workbook.sheetnames)
        self.assertIn("Ad 13 Sal e FGTS", workbook.sheetnames)
        aba = workbook["Ad 13 Sal e FGTS"]
        valores = {r[0]: r[1:3] for r in list(aba.values)[5:]}
        self.assertEqual(valores["F05"], (7117.99, 569.42))
        self.assertNotIn("Mtz", valores)
        self.assertEqual(valores["F116"], (2766.07, 221.28))
        self.assertEqual(len(valores), 16)
        self.assertEqual(valores["TOTAL GERAL"], ('=SUM(B6:B20)', '=SUM(C6:C20)'))

    def iniciar(self, config=None):
        # Diretório controlado pelo teste evita problemas do sandbox com TEMP do Windows.
        pasta = ROOT / "tmp" / "streamlit_tests"
        pasta.mkdir(parents=True, exist_ok=True)
        anterior = tempfile.tempdir
        tempfile.tempdir = str(pasta)
        self.addCleanup(setattr, tempfile, "tempdir", anterior)
        app = AppTest.from_file(str(ROOT / "tests/support/app_pdf.py"), default_timeout=60)
        app.secrets["supabase_url"] = "https://example.supabase.co"
        app.secrets["supabase_key"] = "test-key"
        if config is not None:
            app.session_state["test_db"] = {CHAVE_CONFIG_PDF: config}
        return app.run()

    def test_pdf_sem_csv_salva_recarrega_e_exporta(self):
        app = self.iniciar()
        self.assertFalse(app.exception)
        self.assertNotIn("test_export", app.session_state)
        self.assertEqual(app.tabs[0].label, "Resumo Salários")
        self.assertEqual(len(app.tabs[0].multiselect), 17)
        app.multiselect(key="pdf_v1_salario").set_value(["00050001"])
        app.multiselect(key=f"pdf_v1_{COLUNAS_PDF[0]}").set_value(["00063101"])
        app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual(len(app.tabs), 7)
        banco = app.session_state["test_db"]
        self.assertEqual(list(banco), [CHAVE_CONFIG_PDF])
        self.assertEqual(banco[CHAVE_CONFIG_PDF]["codigos_padrao"], ["00050001"])
        resumo = next(t.value for t in app.dataframe if "C-270.4 - INSS" in t.value.columns)
        self.assertFalse(any("Ad 13" in str(col) for col in resumo.columns))
        self.assertIn("171.255,81", str(resumo.loc["Matriz", "D-266.6 - T. salário"]))
        self.assertIn("16.274,73", str(resumo.loc["Matriz", "C-270.4 - INSS"]))
        self.assertEqual(int(resumo.loc["Filial 25", "Func."]), 51)
        self.assertEqual(app.session_state["test_export_name"], "resumo_salarios_buffon_09-2026.xlsx")
        workbook = openpyxl.load_workbook(io.BytesIO(app.session_state["test_export"]))
        self.assertIn("Resumo Salários", workbook.sheetnames)
        self.assertIn("FGTS e Honorários", workbook.sheetnames)
        aba_codigos = workbook["Códigos dos Cálculos"]
        self.assertEqual(aba_codigos["A1"].value, workbook["Resumo Salários"]["A1"].value)
        self.assertIn("09/2026", aba_codigos["A2"].value)
        auditoria = list(aba_codigos.values)[2:]
        coluna_salario = auditoria[0].index("D-266.6")
        coluna_inss = auditoria[0].index("C-270.4")
        self.assertTrue(any(
            str(row[coluna_salario]).startswith("00050001 - ")
            for row in auditoria[2:]
        ))
        self.assertTrue(any(
            str(row[coluna_inss]).startswith("00063101 - ")
            for row in auditoria[2:]
        ))
        self.assertEqual(len(auditoria[0]), len(COLUNAS_PDF) + 3)
        self.assertIsNone(auditoria[3][coluna_salario])
        self.assertFalse(any(
            "Ad 13" in str(cell.value)
            for row in workbook["Resumo Salários"] for cell in row
        ))
        app.run()
        self.assertEqual(app.multiselect(key="pdf_v1_salario").value, ["00050001"])
        self.assertFalse(app.error)

    def test_codigo_ausente_pode_ser_removido(self):
        config = config_pdf_inicial()
        config["configurado"] = True
        config["prolabore"] = ["00999999"]
        app = self.iniciar(config)
        self.assertEqual(app.multiselect(key="pdf_v1_prolabore").value, ["00999999"])
        app.multiselect(key="pdf_v1_prolabore").set_value([])
        app.button[0].click().run()
        self.assertFalse(app.error)
        self.assertEqual(app.session_state["test_db"][CHAVE_CONFIG_PDF]["prolabore"], [])

    def test_falha_de_gravacao_nao_libera_calculo_inicial(self):
        app = self.iniciar()
        app.session_state["test_fail_save"] = True
        app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertTrue(app.error)
        self.assertNotIn("test_export", app.session_state)

    def test_csv_continua_com_configuracoes_separadas(self):
        app = self.iniciar()
        app.radio(key="fonte_lancamentos").set_value("CSV (padrão antigo)").run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual(len(app.tabs), 5)
        self.assertNotIn(CHAVE_CONFIG_PDF, app.session_state["test_db"])
        self.assertIn("COLUNAS_ESPECIAIS", app.session_state["test_db"])
