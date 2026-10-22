import io
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path

import fitz
import openpyxl
import pandas as pd

from folha_inss import apurar_inss, extrair_descontos, extrair_gps, tabela_inss, escrever_inss

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / 'Resumo GPS Buffon 09 2026 (1).pdf'
ODS = ROOT / 'Buffon Inss Desconto Empregados 09.2026 (1).ods'


class InssTest(unittest.TestCase):
    def xlsx(self, abas):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for linhas in abas:
            aba = workbook.create_sheet()
            for linha in linhas:
                aba.append(linha)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_xlsx_preserva_numeros_textos_zero_e_validacoes(self):
        linhas = [('1001 - Buffon 01', 1234.56), ('1002 - Buffon 02', '2,44'),
                  ('1003 - Buffon 03', 0), ('Totalizador', 1237)]
        self.assertEqual(extrair_descontos(self.xlsx([linhas, list(reversed(linhas))])),
                         {1001: Decimal('1234.56'), 1002: Decimal('2.44'), 1003: Decimal(0)})
        for abas, erro in [([linhas + [linhas[0]]], 'duplicado'),
                           ([linhas, [('1001 - Buffon 01', 1)]], 'divergentes'),
                           ([[('1001 - Buffon 01', None)]], 'inválido'),
                           ([[('1001 - Buffon 01', 1), ('Total', 2)]], 'totalizador')]:
            with self.assertRaisesRegex(ValueError, erro):
                extrair_descontos(self.xlsx(abas))
        pdf = self.pdf('1001 - Buffon 01\nTotais Do Estabelecimento - parte Empresa\nR$ 10,00')
        dados = self.xlsx([[('1001 - Buffon 01', 1.25)]])
        self.assertEqual(tabela_inss(apurar_inss(pdf, dados, '09.2026.xlsx')).iloc[0].tolist(),
                         ['Mtz', 10, 1.25, 11.25])
        with self.assertRaisesRegex(ValueError, 'competência'):
            apurar_inss(pdf, dados, '08.2026.xlsx')

    def ods(self, abas):
        xml = '<root xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">'
        for linhas in abas:
            xml += '<table:table>'
            for nome, valor in linhas:
                xml += f'<table:table-row><table:table-cell>{nome}</table:table-cell><table:table-cell>{valor}</table:table-cell></table:table-row>'
            xml += '</table:table>'
        xml += '</root>'
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as arquivo:
            arquivo.writestr('content.xml', xml)
        return output.getvalue()

    def pdf(self, texto):
        with fitz.open() as doc:
            doc.new_page().insert_text((30, 30), 'Mês/Ano de Apuração: 09/2026\n' + texto)
            return doc.tobytes()

    def test_abas_iguais_nao_duplicam_valores_e_abas_divergentes_falham(self):
        linhas = [('1002 - Buffon 02', '2,00'), ('1001 - Buffon 01', '1,00'), ('Totalizador', '3,00')]
        self.assertEqual(extrair_descontos(self.ods([linhas, list(reversed(linhas))])), {1001: Decimal('1'), 1002: Decimal('2')})
        with self.assertRaisesRegex(ValueError, 'divergentes'):
            extrair_descontos(self.ods([linhas, [('1001 - Buffon 01', '4,00')]]))
        with self.assertRaisesRegex(ValueError, 'duplicado'):
            extrair_descontos(self.ods([linhas + [linhas[0]]]))
        with self.assertRaisesRegex(ValueError, 'totalizador'):
            extrair_descontos(self.ods([[('1001 - Buffon 01', '1,00'), ('Totalizador', '2,00')]]))

    def test_incompleto_ou_mes_divergente_nao_vira_soma_valida(self):
        with self.assertRaisesRegex(ValueError, 'ausente'):
            extrair_gps(self.pdf('1001 - Buffon 01'))
        pdf = self.pdf('1001 - Buffon 01\nTotais Do Estabelecimento - parte Empresa\nR$ 0,00')
        with self.assertRaisesRegex(ValueError, 'correspondência'):
            apurar_inss(pdf, self.ods([[('1002 - Buffon 02', '2,00')]]))
        ods = self.ods([[('1001 - Buffon 01', '0,00')]])
        with self.assertRaisesRegex(ValueError, 'competência'):
            apurar_inss(pdf, ods, 'Desconto 08.2026.ods')
        self.assertEqual(tabela_inss(apurar_inss(pdf, ods)).iloc[0].tolist(), ['Mtz', 0, 0, 0])

    @unittest.skipUnless(PDF.exists() and ODS.exists(), 'Arquivos de exemplo ausentes')
    def test_reais_reconciliam_e_excel_segue_padrao(self):
        dados = apurar_inss(PDF.read_bytes(), ODS.read_bytes(), ODS.name)
        self.assertEqual(len(dados['registros']), 116)
        tabela = tabela_inss(dados)
        self.assertEqual(tabela.iloc[0].tolist(), ['Mtz', 53753.59, 16426.92, 70180.51])
        self.assertEqual(tabela.iloc[1]['Filial'], 'F01')
        self.assertEqual(tabela.iloc[-1].tolist(), ['TOTAL GERAL', 2085120.27, 638699.88, 2723820.15])
        self.assertEqual(dados['avisos'], [])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
            escrever_inss(writer, dados)
        aba = openpyxl.load_workbook(io.BytesIO(output.getvalue()))['Resumo INSS']
        self.assertEqual(aba['A2'].value, 'Com Buffon / INSS - 09/2026')
        self.assertEqual(aba['A5'].value, 'Mtz')
        self.assertEqual(aba['D5'].value, '=B5+C5')
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(output.getvalue()), data_only=True)['Resumo INSS']['D5'].value, 70180.51)
        self.assertEqual(aba['B5'].number_format, '#,##0.00')
        self.assertTrue(aba['D121'].font.bold)
