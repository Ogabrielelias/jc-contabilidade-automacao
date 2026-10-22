import io
import unittest
from decimal import Decimal
from pathlib import Path

import fitz
import openpyxl
import pandas as pd

from folha_provisao_13 import apurar_provisao_13, extrair_provisao_13, ler_saldos, tabela_provisao_13, escrever_provisao_13

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / 'Provisão de 13° Buffon Final  09 - 2026.pdf'
ODS = ROOT / 'prov 13 sal _09-2026 (3) nova.ods'


@unittest.skipUnless(PDF.exists() and ODS.exists(), 'Exemplos locais ausentes')
class ProvisaoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dados = apurar_provisao_13(PDF.read_bytes(), ODS.read_bytes(), ODS.name)

    def test_exemplos_e_cobertura(self):
        self.assertEqual(len(self.dados['registros']), 116)
        tabela = tabela_provisao_13(self.dados)
        self.assertEqual(tabela.iloc[:6, -1].tolist(), [175528.40, 51740.02, 47653.00, 149335.37, 102707.51, 130083.78])
        self.assertEqual(tabela.iloc[5].tolist(), ['F05', 6349.09, 100178.49, 93829.40, 30838.93, 11973.69, 4073.16, 2485.08, 130083.78])
        self.assertEqual(self.dados['registros']['Filial 119']['saldo_planilha'], Decimal(0))
        self.assertEqual(tabela.iloc[-1, 1], 0)
        registro = dict(self.dados['registros']['Filial 119'])
        registro['valores'] = [Decimal(100), *registro['valores'][1:]]
        exemplo = {'registros': {'Filial 119': registro}}
        self.assertEqual(tabela_provisao_13(exemplo).iloc[0, 1], 0)
        registro['saldo_ausente'] = False
        self.assertEqual(tabela_provisao_13(exemplo).iloc[0, 1], -100)
        self.assertEqual(self.dados['registros']['Filial 5']['pagina'], 11)

    def test_excel_formulas_e_proximo_mes(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
            escrever_provisao_13(writer, self.dados)
        wb = openpyxl.load_workbook(io.BytesIO(output.getvalue()))
        ws = wb['Provisão 13º']
        self.assertNotIn('Base Provisão', wb.sheetnames)
        self.assertEqual(ws['E11'].value, '=32268.82-1429.89')
        self.assertEqual(ws['F11'].value, '=9407.35+2566.34')
        self.assertEqual(ws['H11'].value, '=1604.32+880.76')
        self.assertEqual(ws['I11'].value, '=D11+E11+F11-G11-H11')
        self.assertEqual(ws['B11'].value, '=C11-D11')
        valores = openpyxl.load_workbook(io.BytesIO(output.getvalue()), data_only=True)['Provisão 13º']
        self.assertEqual(valores['I11'].value, 130083.78)
        self.assertEqual(valores['C121'].value, 0)
        self.assertEqual(valores['B121'].value, 0)
        saldos = ler_saldos(output.getvalue(), 'anterior.xlsx', '10/2026')
        self.assertEqual(saldos['Filial 5'], Decimal('130083.78'))
        self.assertEqual(saldos['Filial 119'], Decimal('7887.92'))
        with self.assertRaisesRegex(ValueError, 'mês anterior'):
            ler_saldos(output.getvalue(), 'anterior.xlsx', '11/2026')

    def test_rejeita_pdf_incompleto_e_total_inconsistente(self):
        for valor, mensagem in [('130.083,78', 'Total calculado'), ('32.268,82', 'incompleto')]:
            with fitz.open(PDF) as original:
                with fitz.open() as doc:
                    doc.insert_pdf(original, from_page=10, to_page=10)
                    page = doc[0]
                    rect = page.search_for(valor)[0]
                    page.add_redact_annot(rect)
                    page.apply_redactions()
                    if valor == '130.083,78':
                        page.insert_text((rect.x0, rect.y1-2), '130.083,79', fontsize=8)
                    with self.assertRaisesRegex(ValueError, mensagem):
                        extrair_provisao_13(doc.tobytes())


if __name__ == '__main__':
    unittest.main()
