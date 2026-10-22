"""Conciliação do resumo GPS e dos descontos de INSS por estabelecimento."""

import io
import re
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal, InvalidOperation

import fitz
import pandas as pd
from openpyxl import load_workbook

VALOR = r"-?[\d.]+,\d{2}"
UNIDADE = r"(\d{4})\s*-\s*Buffon\s+(\d+)"
NS = {"t": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
      "o": "urn:oasis:names:tc:opendocument:xmlns:office:1.0"}
COLUNAS = ("Parte empresa", "Desconto empregados", "Total INSS")


def numero(texto):
    try:
        return Decimal(texto.replace("R$", "").replace(".", "").replace(",", ".").strip())
    except InvalidOperation as erro:
        raise ValueError(f"Valor de INSS inválido: {texto!r}.") from erro


def extrair_gps(conteudo):
    with fitz.open(stream=conteudo, filetype="pdf") as doc:
        texto = "\n".join(p.get_text() for p in doc)
    periodos = set(re.findall(r"Mês/Ano de Apuração:\s*(\d{2}/\d{4})", texto))
    if len(periodos) != 1:
        raise ValueError("Competência ausente ou divergente no resumo GPS.")
    registros = {}
    cabecalhos = list(re.finditer(r"(?m)^" + UNIDADE + r"\s*$", texto))
    for i, match in enumerate(cabecalhos):
        codigo, filial = map(int, match.groups())
        bloco = texto[match.end():cabecalhos[i+1].start() if i+1 < len(cabecalhos) else len(texto)]
        valores = re.findall(r"Totais Do Estabelecimento\s*-\s*parte Empresa\s*R\$\s*(" + VALOR + ")", bloco, re.I)
        if len(valores) != 1 or codigo in registros or filial < 1 or codigo != 1000 + filial:
            raise ValueError(f"Total ausente, duplicado ou identificação inválida no GPS: {codigo}.")
        registros[codigo] = numero(valores[0])
    if not registros:
        raise ValueError("Nenhum total por estabelecimento encontrado no resumo GPS.")
    consolidado = None
    if "Totais da Empresa para Conferência" in texto:
        final = texto.split("Totais da Empresa para Conferência")[-1]
        valores = [re.search(r"(?m)^" + codigo + r"[^\n]*\n\s*(" + VALOR + ")", final)
                   for codigo in ("113801", "113804", "164601", "117001", "117601", "119101", "119601", "120001")]
        if all(valores):
            consolidado = sum(numero(v[1]) for v in valores)
    return {"competencia": periodos.pop(), "registros": registros, "consolidado": consolidado}


def linhas_descontos(conteudo):
    """Lê as duas primeiras colunas preservando zeros e casas decimais."""
    with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo:
        if 'xl/workbook.xml' in arquivo.namelist():
            workbook = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
            try:
                for aba in workbook:
                    linhas = []
                    for nome, valor in aba.iter_rows(max_col=2, values_only=True):
                        nome = '' if nome is None else str(nome).strip()
                        if valor is None:
                            valor = ''
                        elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
                            valor = str(valor).replace('.', ',')
                        else:
                            valor = str(valor).strip()
                        linhas.append(([nome, valor], 1))
                    yield linhas
            finally:
                workbook.close()
            return
        if 'content.xml' not in arquivo.namelist():
            raise ValueError('Envie uma planilha ODS ou XLSX válida.')
        raiz = ET.fromstring(arquivo.read("content.xml"))
    for aba in raiz.findall(".//t:table", NS):
        linhas = []
        for linha in aba.findall("t:table-row", NS):
            celulas = []
            for celula in linha:
                texto = "".join(celula.itertext()).strip()
                if not texto and celula.get("{" + NS['o'] + "}value") is not None:
                    texto = celula.get("{" + NS['o'] + "}value").replace('.', ',')
                repeticoes = int(celula.get("{" + NS['t'] + "}number-columns-repeated", "1"))
                celulas.extend([texto] * min(repeticoes, 2-len(celulas)))
                if len(celulas) == 2:
                    break
            linhas.append((celulas, int(linha.get("{" + NS['t'] + "}number-rows-repeated", "1"))))
        yield linhas


def extrair_descontos(conteudo):
    abas = []
    for linhas in linhas_descontos(conteudo):
        registros, totais = {}, []
        for celulas, repeticoes in linhas:
            if len(celulas) < 2:
                continue
            nome, valor = celulas
            match = re.fullmatch(UNIDADE, nome, re.I)
            if match:
                codigo, filial = map(int, match.groups())
                if codigo in registros or codigo != 1000 + filial or filial < 1:
                    raise ValueError(f"Estabelecimento duplicado ou inválido na planilha: {nome}.")
                if repeticoes != 1:
                    raise ValueError(f"Estabelecimento repetido na planilha: {nome}.")
                registros[codigo] = numero(valor)
            elif 'buffon' in nome.lower():
                raise ValueError(f"Estabelecimento inválido na planilha: {nome}.")
            elif valor and (not nome or nome.lower().startswith('total')):
                totais.append(numero(valor))
        if registros:
            if any(total != sum(registros.values()) for total in totais):
                raise ValueError("O totalizador da planilha difere da soma dos descontos por filial.")
            abas.append(registros)
    if not abas:
        raise ValueError("Nenhum desconto por estabelecimento encontrado na planilha.")
    if any(aba != abas[0] for aba in abas[1:]):
        raise ValueError("As abas da planilha contêm descontos divergentes. Confira o arquivo.")
    return abas[0]


def apurar_inss(pdf, ods, nome_ods=""):
    gps = extrair_gps(pdf)
    descontos = extrair_descontos(ods)
    periodo = re.search(r"(?<!\d)(0[1-9]|1[0-2])[ ._/-]+(20\d{2})(?!\d)", nome_ods)
    if periodo and '/'.join(periodo.groups()) != gps['competencia']:
        raise ValueError("A competência do nome da planilha difere do resumo GPS.")
    if set(gps['registros']) != set(descontos):
        faltam_gps = sorted(set(descontos)-set(gps['registros']))
        faltam_ods = sorted(set(gps['registros'])-set(descontos))
        raise ValueError(f"Filiais sem correspondência. Ausentes no GPS: {faltam_gps}; na planilha: {faltam_ods}.")
    registros = {codigo: {COLUNAS[0]: empresa, COLUNAS[1]: descontos[codigo],
                          COLUNAS[2]: empresa + descontos[codigo]}
                 for codigo, empresa in sorted(gps['registros'].items())}
    return {'competencia': gps['competencia'], 'registros': registros, 'avisos': []}


def tabela_inss(dados):
    linhas = [{"Filial": "Mtz" if codigo == 1001 else f"F{codigo-1001:02d}",
               **{coluna: float(valores[coluna]) for coluna in COLUNAS}}
              for codigo, valores in sorted(dados['registros'].items())]
    linhas.append({"Filial": "TOTAL GERAL", **{
        coluna: float(sum(r[coluna] for r in dados['registros'].values())) for coluna in COLUNAS}})
    return pd.DataFrame(linhas)


def escrever_inss(writer, dados):
    tabela = tabela_inss(dados)
    tabela.to_excel(writer, sheet_name="Resumo INSS", index=False, startrow=3)
    workbook, worksheet = writer.book, writer.sheets['Resumo INSS']
    titulo = workbook.add_format({'bold': True, 'font_size': 14, 'align': 'left', 'valign': 'vcenter'})
    subtitulo = workbook.add_format({'bold': True, 'font_size': 12, 'align': 'left', 'valign': 'vcenter'})
    moeda = workbook.add_format({'num_format': '#,##0.00', 'align': 'right', 'valign': 'vcenter'})
    total = workbook.add_format({'bold': True, 'num_format': '#,##0.00', 'align': 'right', 'valign': 'vcenter'})
    worksheet.merge_range('A1:G1', 'COMERCIAL BUFFON COMB. E TRANSPORTES LTDA', titulo)
    worksheet.merge_range('A2:G2', f"Com Buffon / INSS - {dados['competencia']}", subtitulo)
    worksheet.set_column('B:D', None, moeda)
    worksheet.set_row(len(tabela)+3, None, total)
    for linha in range(4, len(tabela)+3):
        n = linha+1
        worksheet.write_formula(linha, 3, f'=B{n}+C{n}', moeda, float(tabela.iloc[linha-4, 3]))
    for coluna, letra in enumerate('BCD', 1):
        formula = f'=SUM({letra}5:{letra}{len(tabela)+3})' if len(tabela) > 1 else '=0'
        worksheet.write_formula(len(tabela)+3, coluna, formula, total, float(tabela.iloc[-1, coluna]))
    for i, aviso in enumerate(dados['avisos'], len(tabela)+5):
        worksheet.merge_range(i, 0, i+1, 6, aviso, workbook.add_format({'text_wrap': True, 'valign': 'top'}))
    worksheet.autofit()
    return worksheet
